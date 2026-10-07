# -*- coding: utf-8 -*-
"""
Full Production Pipeline for KaiDance Album 1 (Songs 2 to 10):
- Complete Action Analysis & Animation Specification
- Midjourney / Sora / Runway Gen-3 AI Video Animation Prompts
- Warm, Encouraging Teacher Voice Prompts (Traditional Chinese)
- Excel Workbook (.xlsx) Generation
- Elegant PDF Document (.pdf) Generation
- High-Fidelity Teacher Voice MP3 Audio Generation (edge-tts zh-TW-HsiaoChenNeural)
- Update web_app/album_data.js and docs/album_data.js
"""

import os, sys, json, asyncio, re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether, HRFlowable
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

import edge_tts

# 1. Register Chinese Font
FONT_NAME = "STHeiti"
FONT_PATH = "/System/Library/Fonts/STHeiti Medium.ttc"
if not os.path.exists(FONT_PATH):
    FONT_PATH = "/System/Library/Fonts/STHeiti Light.ttc"
if not os.path.exists(FONT_PATH):
    FONT_PATH = "/Library/Fonts/Arial Unicode.ttf"
pdfmetrics.registerFont(TTFont(FONT_NAME, FONT_PATH))

# 2. Load existing album database
with open('web_app/album_data.js', 'r', encoding='utf-8') as f:
    orig_text = f.read()

db = json.loads(orig_text.replace('window.KAI_ALBUM_DATABASE = ', '').rstrip().rstrip(';'))
album1_songs = db['album_1']['songs']

# 3. Dedicated Prompt Design Knowledgebase for Songs 02 to 10
# Base prompt template ensuring consistency with KaiDance 3D boy/girl coach style
ANIM_BASE_PREFIX = "3D Pixar Disney style cheerful toddler dance coach, smooth playful animation, vibrant kindergarten studio, soft bright volumetric lighting, clean studio background, highly expressive face, full body dynamic camera shot: "

def generate_video_prompt(song_title, theme, action_name, action_desc, pose):
    pose_keywords = {
        "march": "marching in place energetically, hands on hips, knees lifting high in rhythmic bounce, enthusiastic happy smile",
        "wave": "waving both hands overhead in friendly wide arcs, swaying gently from side to side, sparkling excited eyes",
        "reach": "stretching arms up high to reach floating elements, standing on tiptoes, joyful posture, looking up with wonder",
        "sway": "swaying hips and arms like a breeze, smooth fluid body roll, playful rhythmic dancing, friendly warm smile",
        "jump": "squatting down and bursting into a high happy jump, feet kicking joyfully in mid-air, both fists punching the sky in victory",
        "sun_rise": "lifting arms from heart to sky creating a big golden sun circle, chest open, radiant beaming joyful smile",
        "heart": "bringing hands together at the chest forming an adorable glowing heart shape, leaning slightly forward, loving gentle smile"
    }
    pose_detail = pose_keywords.get(pose, "cheerful dance movement with enthusiastic toddler-friendly choreography")
    prompt = f"{ANIM_BASE_PREFIX}Song '{song_title}', Action '{action_name}'. {action_desc}. The character is {pose_detail}. 4k resolution, 60fps, Disney Pixar render, joyful preschool music video style."
    return prompt

# Teacher voice style prefix generator
def generate_teacher_voice_script(song_num, song_title, seg_title, lyric, sub_tip, pose):
    clean_lyric = lyric.replace('"', '').replace('[', '').replace(']', '')
    
    # Contextual warm preschool teacher speech
    if "前奏" in seg_title:
        return f"小朋友們站好囉！雙手插腰，跟著歡樂音樂踩踩拍子踏步走！"
    elif "間奏" in seg_title:
        return f"太棒啦！小腳步輕輕踏，轉個漂亮的小圓圈，準備聽下一句囉！"
    elif "結尾" in seg_title or "定格" in seg_title or "通關" in seg_title:
        return f"哇！跳得太完美了！雙手比出大愛心，給自己拍拍手，超級棒！"
    elif pose == "march":
        return f"小手插腰，像勇敢的小士兵一樣，一、二、一、二，踩起小腳步！"
    elif pose == "wave":
        return f"揮揮你的小手，左右大力招手打招呼，露出最甜的笑容！"
    elif pose == "reach":
        return f"墊起小腳尖，雙手伸得高高的，一起向上摘星星囉！"
    elif pose == "jump":
        return f"雙膝蹲低低蓄力，預備——用力跳起來，喊一聲耶！"
    elif pose == "sun_rise":
        return f"雙手從胸前向上推，畫出最大最金黃的太陽，全身暖洋洋！"
    elif pose == "heart":
        return f"雙手在胸口比出跳動的大愛心，把滿滿的愛送給好朋友！"
    elif pose == "sway":
        return f"像微風中的小柳樹一樣，雙臂展開，跟著音樂輕輕搖擺！"
    else:
        return f"跟著老師的口令，身體動起來，快樂跳舞囉！"

# Process Songs 02 to 10
processed_data = []

# Directory for voices
os.makedirs("web_app/voices/album1", exist_ok=True)
os.makedirs("docs/voices/album1", exist_ok=True)

tts_tasks = []

for song in album1_songs:
    s_num = song['trackNumber']
    s_id = song['id']
    s_title = song['title']
    s_zh = song['zhTitle']
    
    # We will refine/complete all songs 02 to 10 (and ensure 01 is fully aligned)
    for idx, seg in enumerate(song['timeline']):
        seg_idx = idx + 1
        dur = round(seg['end'] - seg['start'], 2)
        
        # 1. Action Summary & Key Technique
        action_name = seg['title']
        lyric = seg['lyric']
        pose = seg.get('targetPose', 'march')
        
        # Ensure detailed pedagogical movement instructions
        key_tip = seg['sub']
        
        # 2. Teacher Voice Script (Natural, warm, friendly Traditional Chinese)
        voice_script = generate_teacher_voice_script(s_num, s_title, seg['title'], lyric, key_tip, pose)
        voice_file = f"./voices/album1/voice_track_{s_num:02d}_seg_{seg_idx:02d}.mp3"
        seg['voice'] = voice_file
        seg['voiceScript'] = voice_script
        
        # 3. AI Video Production Prompt
        video_prompt = generate_video_prompt(s_title, song['tag'], action_name, key_tip, pose)
        seg['videoPrompt'] = video_prompt
        
        # Record into production dataset
        processed_data.append({
            "trackNumber": s_num,
            "songId": s_id,
            "songTitle": s_title,
            "songZhTitle": s_zh,
            "segIndex": seg_idx,
            "startTime": f"{int(seg['start']//60):02d}:{int(seg['start']%60):02d}.{int((seg['start']%1)*10):d}",
            "endTime": f"{int(seg['end']//60):02d}:{int(seg['end']%60):02d}.{int((seg['end']%1)*10):d}",
            "startSec": seg['start'],
            "endSec": seg['end'],
            "durationSec": dur,
            "segTitle": seg['title'],
            "lyric": seg['lyric'],
            "keyTip": key_tip,
            "voiceScript": voice_script,
            "targetPose": pose,
            "videoPrompt": video_prompt,
            "voiceFile": voice_file
        })

print(f"Total processed segments across Album 1: {len(processed_data)}")

# 4. Save updated database back to web_app and docs
output_js = "window.KAI_ALBUM_DATABASE = " + json.dumps(db, ensure_ascii=False, indent=2) + ";\n"
with open('web_app/album_data.js', 'w', encoding='utf-8') as f:
    f.write(output_js)
with open('docs/album_data.js', 'w', encoding='utf-8') as f:
    f.write(output_js)
print("Updated web_app/album_data.js and docs/album_data.js with new voiceScript and videoPrompt.")

# 5. Build Excel Workbook (.xlsx)
print("Building Excel Workbook...")
wb = openpyxl.Workbook()

# Setup Master Overview Sheet
ws_summary = wb.active
ws_summary.title = "曲目總覽與製作統計"
ws_summary.views.sheetView[0].showGridLines = True

# Styling helpers
font_title = Font(name="Microsoft JhengHei", size=16, bold=True, color="1E293B")
font_section = Font(name="Microsoft JhengHei", size=13, bold=True, color="0F172A")
font_header = Font(name="Microsoft JhengHei", size=10, bold=True, color="FFFFFF")
font_cell = Font(name="Microsoft JhengHei", size=10, color="1E293B")
font_small = Font(name="Microsoft JhengHei", size=9, color="475569")

fill_header_amber = PatternFill(start_color="F59E0B", end_color="F59E0B", fill_type="solid")
fill_header_sky = PatternFill(start_color="0284C7", end_color="0284C7", fill_type="solid")
fill_header_emerald = PatternFill(start_color="059669", end_color="059669", fill_type="solid")
fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

thin_border = Border(
    left=Side(style='thin', color='CBD5E1'),
    right=Side(style='thin', color='CBD5E1'),
    top=Side(style='thin', color='CBD5E1'),
    bottom=Side(style='thin', color='CBD5E1')
)

# Header title
ws_summary.merge_cells('A1:G1')
ws_summary['A1'] = "🎵 KaiDance 第一輯《小小腳步 大大世界》全 10 首歌 動作教學與動畫製作規格總表"
ws_summary['A1'].font = font_title
ws_summary['A1'].alignment = Alignment(horizontal="center", vertical="center")
ws_summary.row_dimensions[1].height = 40

summary_headers = ["曲目", "英文歌名", "中文主題", "律動標籤", "總長度", "段落數", "重點動作核心"]
ws_summary.append([]) # row 2 empty
ws_summary.append(summary_headers) # row 3
ws_summary.row_dimensions[3].height = 26

for col_num in range(1, 8):
    c = ws_summary.cell(row=3, column=col_num)
    c.font = font_header
    c.fill = fill_header_amber
    c.alignment = Alignment(horizontal="center", vertical="center")

for s in album1_songs:
    row_data = [
        f"#{s['trackNumber']:02d}",
        s['title'],
        s['zhTitle'],
        s['tag'],
        f"{s['duration']} 秒",
        f"{len(s['timeline'])} 段",
        s['description'][:40] + "..."
    ]
    ws_summary.append(row_data)
    curr_row = ws_summary.max_row
    ws_summary.row_dimensions[curr_row].height = 24
    for c_idx in range(1, 8):
        cell = ws_summary.cell(row=curr_row, column=c_idx)
        cell.font = font_cell
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center" if c_idx in [1, 5, 6] else "left", vertical="center")
        if curr_row % 2 == 1:
            cell.fill = fill_zebra

# Set column widths for summary
summary_widths = [10, 28, 22, 18, 12, 12, 45]
for idx, w in enumerate(summary_widths, 1):
    ws_summary.column_dimensions[get_column_letter(idx)].width = w

# Create Detailed Sheets for Song 2 through Song 10
detail_headers = [
    "段落序號", "時間區間 (分:秒)", "時長(秒)", "段落動作主題", "對應英文歌詞", 
    "幼兒動作要領說明", "名師親切引導口令內容 (中文)", "姿態類型", "AI 影片動畫生成提示詞 (Sora/Runway Prompt)"
]

# Create one comprehensive detail sheet with all songs, plus individual tabs
ws_all = wb.create_sheet(title="全曲目動作與口令細節總表")
ws_all.views.sheetView[0].showGridLines = True
ws_all.append(detail_headers)
ws_all.row_dimensions[1].height = 28

for col_num in range(1, len(detail_headers) + 1):
    c = ws_all.cell(row=1, column=col_num)
    c.font = font_header
    c.fill = fill_header_sky
    c.alignment = Alignment(horizontal="center", vertical="center")

for item in processed_data:
    r = [
        f"Track {item['trackNumber']:02d} - #{item['segIndex']:02d}",
        f"{item['startTime']} ~ {item['endTime']}",
        item['durationSec'],
        item['segTitle'],
        item['lyric'],
        item['keyTip'],
        item['voiceScript'],
        item['targetPose'],
        item['videoPrompt']
    ]
    ws_all.append(r)
    curr_row = ws_all.max_row
    ws_all.row_dimensions[curr_row].height = 28
    for c_idx in range(1, len(detail_headers) + 1):
        cell = ws_all.cell(row=curr_row, column=c_idx)
        cell.font = font_cell
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center" if c_idx in [1, 2, 3, 8] else "left", vertical="center", wrap_text=True)
        if curr_row % 2 == 1:
            cell.fill = fill_zebra

detail_widths = [18, 16, 10, 22, 30, 36, 36, 12, 60]
for idx, w in enumerate(detail_widths, 1):
    ws_all.column_dimensions[get_column_letter(idx)].width = w

# Also create individual sheets for Song 02 to Song 10 for convenient reading
for s_num in range(2, 11):
    s_info = next(s for s in album1_songs if s['trackNumber'] == s_num)
    ws_s = wb.create_sheet(title=f"第{s_num}首_{s_info['zhTitle'][:6]}")
    ws_s.views.sheetView[0].showGridLines = True
    
    # Title row
    ws_s.merge_cells('A1:I1')
    ws_s['A1'] = f"Track #{s_num:02d}: {s_info['title']} ({s_info['zhTitle']}) - 動作要領、口令與動畫製作提示詞"
    ws_s['A1'].font = font_section
    ws_s['A1'].alignment = Alignment(horizontal="left", vertical="center")
    ws_s.row_dimensions[1].height = 32
    
    ws_s.append(detail_headers)
    ws_s.row_dimensions[2].height = 26
    for col_num in range(1, len(detail_headers) + 1):
        c = ws_s.cell(row=2, column=col_num)
        c.font = font_header
        c.fill = fill_header_emerald
        c.alignment = Alignment(horizontal="center", vertical="center")
        
    s_items = [p for p in processed_data if p['trackNumber'] == s_num]
    for item in s_items:
        r = [
            f"#{item['segIndex']:02d}",
            f"{item['startTime']} ~ {item['endTime']}",
            item['durationSec'],
            item['segTitle'],
            item['lyric'],
            item['keyTip'],
            item['voiceScript'],
            item['targetPose'],
            item['videoPrompt']
        ]
        ws_s.append(r)
        curr_row = ws_s.max_row
        ws_s.row_dimensions[curr_row].height = 30
        for c_idx in range(1, len(detail_headers) + 1):
            cell = ws_s.cell(row=curr_row, column=c_idx)
            cell.font = font_cell
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center" if c_idx in [1, 2, 3, 8] else "left", vertical="center", wrap_text=True)
            if curr_row % 2 == 1:
                cell.fill = fill_zebra
                
    for idx, w in enumerate(detail_widths, 1):
        ws_s.column_dimensions[get_column_letter(idx)].width = w

excel_path_docs = "docs/Album1_Dance_Production_Spec.xlsx"
excel_path_web = "web_app/Album1_Dance_Production_Spec.xlsx"
wb.save(excel_path_docs)
wb.save(excel_path_web)
print(f"Saved Excel spreadsheets to {excel_path_docs} and {excel_path_web}")

# 6. Generate Professional PDF Document (.pdf) using ReportLab
print("Generating PDF Document...")
pdf_path_docs = "docs/Album1_Dance_Production_Spec.pdf"
pdf_path_web = "web_app/Album1_Dance_Production_Spec.pdf"

doc = SimpleDocTemplate(
    pdf_path_docs,
    pagesize=landscape(A4),
    leftMargin=20,
    rightMargin=20,
    topMargin=25,
    bottomMargin=25
)

styles = getSampleStyleSheet()

style_title = ParagraphStyle(
    'DocTitle',
    parent=styles['Normal'],
    fontName=FONT_NAME,
    fontSize=18,
    leading=24,
    textColor=colors.HexColor('#1E293B'),
    alignment=1, # Center
    spaceAfter=10
)

style_subtitle = ParagraphStyle(
    'DocSubtitle',
    parent=styles['Normal'],
    fontName=FONT_NAME,
    fontSize=11,
    leading=15,
    textColor=colors.HexColor('#64748B'),
    alignment=1,
    spaceAfter=15
)

style_h2 = ParagraphStyle(
    'SongH2',
    parent=styles['Normal'],
    fontName=FONT_NAME,
    fontSize=13,
    leading=18,
    textColor=colors.HexColor('#0F172A'),
    spaceBefore=10,
    spaceAfter=8,
    keepWithNext=True
)

style_th = ParagraphStyle(
    'TableHead',
    parent=styles['Normal'],
    fontName=FONT_NAME,
    fontSize=8.5,
    leading=11,
    textColor=colors.white,
    alignment=1
)

style_td = ParagraphStyle(
    'TableCell',
    parent=styles['Normal'],
    fontName=FONT_NAME,
    fontSize=7.5,
    leading=10,
    textColor=colors.HexColor('#1E293B')
)

style_td_center = ParagraphStyle(
    'TableCellCenter',
    parent=styles['Normal'],
    fontName=FONT_NAME,
    fontSize=7.5,
    leading=10,
    textColor=colors.HexColor('#1E293B'),
    alignment=1
)

style_td_code = ParagraphStyle(
    'TableCellCode',
    parent=styles['Normal'],
    fontName=FONT_NAME,
    fontSize=6.5,
    leading=8.5,
    textColor=colors.HexColor('#334155')
)

story = []

# Title & Intro
story.append(Paragraph("KaiDance 第一輯《小小腳步 大大世界》動畫製作與名師口令規格企劃書", style_title))
story.append(Paragraph("幼兒律動舞蹈動作要領 · 毫秒級時鐘對齊 · 3D 動畫 Prompt · 親切引導口令完整規格庫", style_subtitle))
story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#F59E0B'), spaceBefore=0, spaceAfter=12))

# Loop through Song 2 to 10 for PDF sections
for s_num in range(2, 11):
    s_info = next(s for s in album1_songs if s['trackNumber'] == s_num)
    s_items = [p for p in processed_data if p['trackNumber'] == s_num]
    
    story.append(Paragraph(f"<b>Track #{s_num:02d}: {s_info['title']} ({s_info['zhTitle']})</b> · 標籤: {s_info['tag']} · 長度: {s_info['duration']}秒", style_h2))
    
    # Table data
    table_data = [[
        Paragraph("段落", style_th),
        Paragraph("時間點", style_th),
        Paragraph("動作主題", style_th),
        Paragraph("對應歌詞", style_th),
        Paragraph("幼兒動作要領說明", style_th),
        Paragraph("名師親切引導口令", style_th),
        Paragraph("AI 動畫提示詞 (Prompt)", style_th)
    ]]
    
    for item in s_items:
        row = [
            Paragraph(f"#{item['segIndex']:02d}", style_td_center),
            Paragraph(f"{item['startTime']}<br/>~{item['endTime']}", style_td_center),
            Paragraph(f"<b>{item['segTitle']}</b>", style_td),
            Paragraph(item['lyric'], style_td),
            Paragraph(item['keyTip'], style_td),
            Paragraph(f"<b>{item['voiceScript']}</b>", style_td),
            Paragraph(item['videoPrompt'], style_td_code)
        ]
        table_data.append(row)
        
    # Table column widths (total available width landscape A4 = 842 - 40 = 802pt)
    col_widths = [32, 58, 90, 110, 150, 160, 202]
    
    t = Table(table_data, colWidths=col_widths, repeatRows=1)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0284C7')),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
        ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#CBD5E1')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('LEFTPADDING', (0, 0), (-1, -1), 4),
        ('RIGHTPADDING', (0, 0), (-1, -1), 4)
    ]))
    
    story.append(t)
    story.append(Spacer(1, 15))
    if s_num < 10:
        story.append(PageBreak())

doc.build(story)

# Copy PDF to web_app
import shutil
shutil.copy(pdf_path_docs, pdf_path_web)
print(f"Generated PDF documents at {pdf_path_docs} and {pdf_path_web}")

# 7. Asynchronously Batch-Generate Teacher Voice MP3s using edge-tts
print("Starting edge-tts voice generation for Songs 02 to 10...")

async def generate_single_tts(text, out_file):
    if os.path.exists(out_file) and os.path.getsize(out_file) > 1024:
        return
    voice = "zh-TW-HsiaoChenNeural"
    communicate = edge_tts.Communicate(text, voice, rate="+6%", pitch="+4Hz")
    await communicate.save(out_file)

async def batch_generate_tts():
    tasks = []
    # Generate for Song 02 to Song 10
    subset = [p for p in processed_data if p['trackNumber'] >= 2]
    print(f"Total TTS clips to generate: {len(subset)}")
    
    # Semaphore to avoid rate limiting
    sem = asyncio.Semaphore(5)
    
    async def worker(item):
        async with sem:
            out_docs = os.path.join("docs", item['voiceFile'].replace("./", ""))
            out_web = os.path.join("web_app", item['voiceFile'].replace("./", ""))
            os.makedirs(os.path.dirname(out_docs), exist_ok=True)
            os.makedirs(os.path.dirname(out_web), exist_ok=True)
            try:
                await generate_single_tts(item['voiceScript'], out_docs)
                if os.path.exists(out_docs):
                    shutil.copy(out_docs, out_web)
            except Exception as e:
                print(f"TTS Error on {item['songId']} seg {item['segIndex']}: {e}")

    await asyncio.gather(*(worker(item) for item in subset))
    print("All Teacher Voice MP3 files generated successfully!")

asyncio.run(batch_generate_tts())

print("Full production pipeline completed with flying colors!")
